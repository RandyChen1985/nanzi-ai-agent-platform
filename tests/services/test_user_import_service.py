"""Excel 导入用户的解析 / 角色匹配 / 密码规则（纯逻辑）。"""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from app.services.auth_service import AuthService
from app.services.user_import_service import (
    MAX_IMPORT_ROWS,
    UserImportError,
    build_template_workbook,
    cell_to_text,
    find_default_role,
    generate_initial_password,
    match_role,
    parse_import_workbook,
    validate_fixed_password,
    validate_user_name,
)

pytestmark = pytest.mark.no_infrastructure


def _roles():
    return [
        SimpleNamespace(id=1, code="default", name="普通用户", description="默认角色"),
        SimpleNamespace(id=2, code="admin", name="系统管理员", description="管理员"),
        SimpleNamespace(id=3, code="analyst", name="分析师", description=""),
    ]


def _workbook_bytes(rows, sheet_title="用户导入"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- 模板 ---


def test_template_has_four_columns_and_role_guide():
    content = build_template_workbook(_roles())
    workbook = load_workbook(io.BytesIO(content))

    assert workbook.sheetnames == ["用户导入", "填写说明"]
    data_sheet = workbook["用户导入"]
    assert [data_sheet.cell(row=1, column=column).value for column in (1, 2, 3, 4)] == [
        "账号*",
        "姓名",
        "邮箱*",
        "角色",
    ]
    assert data_sheet.auto_filter.ref == "A1:D1"
    # 账号列必须按文本格式，否则「00123」会被 Excel 吃掉前导零
    assert data_sheet.cell(row=2, column=1).number_format == "@"

    guide = workbook["填写说明"]
    values = [
        cell.value
        for row in guide.iter_rows(min_row=1, max_row=guide.max_row, max_col=3)
        for cell in row
    ]
    joined = "\n".join(str(value or "") for value in values)
    assert "初始密码" in joined
    assert "姓名" in joined
    assert "普通用户" in joined
    assert "analyst" in joined
    assert "分析师" in joined


def test_template_is_parseable_even_without_data_rows():
    """空模板用于下载，不应被当成可导入文件；解析应给出明确提示。"""

    content = build_template_workbook(_roles())
    with pytest.raises(UserImportError, match="没有可导入的数据行"):
        parse_import_workbook(content)


# --- 解析 ---


def test_parse_reads_rows_and_skips_empty_lines():
    content = _workbook_bytes(
        [
            ["账号", "姓名", "邮箱", "角色"],
            ["alice", "爱丽丝", "Alice@Example.com", "default"],
            [None, None, None, None],
            ["bob", "鲍勃", "bob@example.com", "分析师"],
        ]
    )
    rows = parse_import_workbook(content)

    assert [
        (row.row_number, row.user_name, row.real_name, row.email, row.role_input)
        for row in rows
    ] == [
        (2, "alice", "爱丽丝", "Alice@Example.com", "default"),
        (4, "bob", "鲍勃", "bob@example.com", "分析师"),
    ]


def test_parse_accepts_header_aliases_and_star_suffix():
    content = _workbook_bytes(
        [
            ["序号"],
            ["用户名*", "真实姓名", "电子邮件", "业务角色"],
            ["carol", "卡罗尔", "carol@example.com", "admin"],
        ]
    )
    rows = parse_import_workbook(content)

    assert len(rows) == 1
    assert rows[0].user_name == "carol"
    assert rows[0].real_name == "卡罗尔"
    assert rows[0].email == "carol@example.com"
    assert rows[0].role_input == "admin"


def test_parse_converts_numeric_account_without_decimal_suffix():
    content = _workbook_bytes(
        [
            ["账号", "邮箱", "角色"],
            [30049, "u30049@example.com", "default"],
        ]
    )
    rows = parse_import_workbook(content)

    assert rows[0].user_name == "30049"


def test_parse_rejects_workbook_without_email_column():
    content = _workbook_bytes([["账号", "角色"], ["dave", "default"]])

    with pytest.raises(UserImportError, match="邮箱"):
        parse_import_workbook(content)


def test_parse_rejects_workbook_without_account_column():
    content = _workbook_bytes([["邮箱", "角色"], ["dave@example.com", "default"]])

    with pytest.raises(UserImportError, match="账号"):
        parse_import_workbook(content)


def test_parse_rejects_non_xlsx_bytes():
    with pytest.raises(UserImportError, match="无法解析"):
        parse_import_workbook(b"not-an-excel-file")


def test_parse_rejects_more_than_row_limit():
    rows = [["账号", "邮箱", "角色"]]
    rows.extend([f"user{i}", f"user{i}@example.com", "default"] for i in range(MAX_IMPORT_ROWS + 1))
    content = _workbook_bytes(rows)

    with pytest.raises(UserImportError, match=str(MAX_IMPORT_ROWS)):
        parse_import_workbook(content)


# --- 角色匹配 ---


def test_match_role_by_code_or_name_case_insensitive():
    roles = _roles()

    assert match_role(" DEFAULT ", roles).id == 1
    assert match_role("分析师", roles).id == 3
    assert match_role("ANALYST", roles).id == 3
    assert match_role("不存在的角色", roles) is None
    assert match_role("", roles) is None


def test_match_role_prefers_code_when_name_collides():
    roles = [
        SimpleNamespace(id=10, code="first", name="同名角色"),
        SimpleNamespace(id=11, code="second", name="同名角色"),
    ]

    assert match_role("同名角色", roles).id == 10
    assert match_role("second", roles).id == 11


def test_find_default_role_matches_code_only():
    roles = _roles()

    assert find_default_role(roles).id == 1
    assert find_default_role([SimpleNamespace(id=9, code="other", name="default")]) is None


# --- 账号与密码 ---


def test_validate_user_name_rules():
    assert validate_user_name("alice") is None
    assert validate_user_name("  alice  ") is None
    assert validate_user_name("") == "账号不能为空"
    assert validate_user_name("a" * 51) == "账号长度不能超过 50 个字符"
    assert validate_user_name("a lice") == "账号不能包含空格"


@pytest.mark.parametrize("user_name", ["alice", "bob", "中文账号"])
def test_initial_password_meets_complexity_and_excludes_username(user_name):
    for _ in range(20):
        password = generate_initial_password(user_name)
        valid, message = AuthService.validate_password_complexity(password, username=user_name)
        assert valid, message
        assert len(password) == 12
        assert not any(char.isspace() for char in password)
        assert not any(char in password for char in ',";')


def test_cell_to_text_handles_numeric_and_boolean_cells():
    assert cell_to_text(None) == ""
    assert cell_to_text(30049.0) == "30049"
    assert cell_to_text(30049.5) == "30049.5"
    assert cell_to_text(True) == "TRUE"
    assert cell_to_text("  alice  ") == "alice"


class _FakeSession:
    """只实现建号路径用到的方法，用来验证「一次提交」而不是连库。"""

    def __init__(self):
        self.added = []
        self.committed = False
        self.rolled_back = False

    def add(self, obj):
        self.added.append(obj)

    def add_all(self, objs):
        self.added.extend(objs)

    async def flush(self):
        for obj in self.added:
            if getattr(obj, "user_name", None) and getattr(obj, "id", None) is None:
                obj.id = 42

    async def commit(self):
        self.committed = True

    async def rollback(self):
        self.rolled_back = True

    async def close(self):
        pass


@pytest.mark.asyncio
async def test_generate_api_key_commits_password_and_roles_together(monkeypatch):
    from app.models.permission import UserRoleRelation
    from app.models.user import User
    from app.services.auth_service import AuthService

    session = _FakeSession()

    class _Manager:
        def generate_api_key(self):
            return "raw-key", "encrypted-key", "hashed-key"

    monkeypatch.setattr(
        "app.services.auth_service.get_api_key_manager", lambda: _Manager()
    )
    monkeypatch.setattr(
        AuthService, "get_password_hash", staticmethod(lambda password: f"hashed:{password}")
    )

    await AuthService.generate_api_key(
        "alice",
        email="alice@example.com",
        db=session,
        password="Abc!12345678",
        business_role_ids=[7],
    )

    users = [obj for obj in session.added if isinstance(obj, User)]
    relations = [obj for obj in session.added if isinstance(obj, UserRoleRelation)]

    assert len(users) == 1
    assert users[0].password_hash == "hashed:Abc!12345678"
    assert users[0].password_updated_at is not None
    assert len(relations) == 1
    assert relations[0].user_id == 42
    assert relations[0].role_id == 7
    assert session.committed is True
    assert session.rolled_back is False


def test_validate_fixed_password_checks_strength_and_contains_account_name():
    assert validate_fixed_password("Abc!12345678", ["alice"]) is None

    assert "长度" in (validate_fixed_password("Abc!1", []) or "")
    assert "3 种" in (validate_fixed_password("abcdefgh", []) or "")

    issue = validate_fixed_password("Alice!2345678", ["alice"])
    assert issue and "包含账号名" in issue, "固定密码复用多个账号，绝不能再包含其中某个账号名"

    # 大写/小写都按同一口径比较；长度 < 3 的账号名不参与，避免误伤
    assert "包含账号名" in (validate_fixed_password("xxALICE!2345", ["Alice"]) or "")
    assert validate_fixed_password("Abc!12345678", ["ab"]) is None


def test_account_created_mail_intentionally_carries_initial_password():
    """这是唯一允许携带凭据的邮件文案，且必须给出改密与防转发提示。"""

    from app.services.email_delivery_service import build_account_created_mail

    subject, body = build_account_created_mail("alice", "Abc!12345678", "NanZi")

    assert "NanZi" in subject and "账号已创建" in subject
    assert "alice" in body
    assert "Abc!12345678" in body
    assert "修改密码" in body
    assert "请勿转发" in body


def test_account_created_mail_includes_platform_url_before_credentials():
    """用户拿到邮件后要能直接知道去哪登录，地址必须在账号密码之前。"""

    from app.services.email_delivery_service import build_account_created_mail

    subject, body = build_account_created_mail(
        "alice", "Abc!12345678", "NanZi", platform_url="https://nanzi.example.com/"
    )

    assert "https://nanzi.example.com" in body
    # 尾斜杠被规整，避免出现 https://host//login
    assert "https://nanzi.example.com/" not in body
    assert body.index("平台地址") < body.index("初始密码")
    assert "修改密码" in body


def test_account_created_mail_omits_platform_line_when_url_is_blank():
    """未配置对外地址时不要渲染一行空的「平台地址：」，那比没有更让人困惑。"""

    from app.services.email_delivery_service import build_account_created_mail

    for blank in ("", "   ", None):
        subject, body = build_account_created_mail(
            "alice", "Abc!12345678", "NanZi", platform_url=blank
        )
        assert "平台地址" not in body
        assert "Abc!12345678" in body
